"""书架服务：上传 → 解析 → （后台）向量化。分析产物另见 doc_analysis。"""
from __future__ import annotations

import re
import uuid
from pathlib import Path
from typing import Any, Optional

from sqlmodel import Session, select

from app.config import DATA_DIR, DOCS_DIR
from app.parser import documents as parser
from app.storage.db import engine
from app.storage.models import Category, Document

ALLOWED_EXTS = {"pdf", "docx", "pptx", "md", "txt", "markdown"}

_FMT_ICON = {"pdf": "📕", "docx": "📘", "pptx": "📙", "md": "📗", "txt": "📄"}


def list_categories() -> list[dict]:
    with Session(engine) as s:
        cats = s.exec(select(Category).order_by(Category.position)).all()
        docs = s.exec(select(Document)).all()
        return [{
            "id": c.id, "name": c.name, "icon": c.icon,
            "count": sum(1 for d in docs if d.category_id == c.id),
        } for c in cats]


def create_category(name: str, icon: str = "📂") -> dict:
    with Session(engine) as s:
        cat = Category(name=name.strip() or "新分类", icon=icon or "📂",
                       position=len(s.exec(select(Category)).all()))
        s.add(cat)
        s.commit()
        s.refresh(cat)
        return {"id": cat.id, "name": cat.name, "icon": cat.icon, "count": 0}


def rename_category(cat_id: int, *, name: Optional[str] = None, icon: Optional[str] = None) -> dict:
    with Session(engine) as s:
        cat = s.get(Category, cat_id)
        if not cat:
            raise KeyError("分类不存在")
        if name and name.strip():
            cat.name = name.strip()
        if icon:
            cat.icon = icon
        s.add(cat)
        s.commit()
        s.refresh(cat)
        docs = s.exec(select(Document)).all()
        return {"id": cat.id, "name": cat.name, "icon": cat.icon,
                "count": sum(1 for d in docs if d.category_id == cat.id)}


def reorder_categories(ids: list[int]) -> None:
    with Session(engine) as s:
        for pos, cid in enumerate(ids):
            cat = s.get(Category, cid)
            if cat:
                cat.position = pos
                s.add(cat)
        s.commit()


def delete_category(cat_id: int) -> None:
    with Session(engine) as s:
        cat = s.get(Category, cat_id)
        if not cat:
            raise KeyError("分类不存在")
        for d in s.exec(select(Document).where(Document.category_id == cat_id)).all():
            d.category_id = None
            s.add(d)
        s.delete(cat)
        s.commit()


def save_upload(filename: str, data: bytes, *, category_id: Optional[int] = None,
                title: Optional[str] = None) -> dict:
    """保存上传文件 → 解析文本 → 入库 → 入队向量索引。"""
    ext = Path(filename).suffix.lstrip(".").lower()
    ext = {"markdown": "md"}.get(ext, ext)
    if ext not in ALLOWED_EXTS:
        raise ValueError(f"不支持的格式 .{ext}（支持 PDF / Word / PPT / Markdown / TXT）")
    doc_id = uuid.uuid4().hex[:8]
    safe_name = re.sub(r"[\\/:*?\"<>|]", "_", Path(filename).name)
    stored = DOCS_DIR / f"{doc_id}-{safe_name}"
    stored.write_bytes(data)

    from app.services import rag_service

    try:
        parsed = parser.parse_any(stored, ext)
        pages = parsed["pages"]
        text = parsed["text"]
    except ValueError:
        stored.unlink(missing_ok=True)
        raise
    except Exception as e:  # noqa: BLE001
        stored.unlink(missing_ok=True)
        raise ValueError(f"解析失败：{e}") from e

    with Session(engine) as s:
        doc = Document(
            title=(title or Path(filename).stem).strip() or Path(filename).stem,
            filename=safe_name,
            path=str(stored.relative_to(DATA_DIR)),
            fmt=ext,
            size_bytes=len(data),
            pages=pages,
            category_id=category_id,
            text_content=text,
            text_chars=len(text),
            ocr_used=bool(parsed.get("ocr_used")),
        )
        s.add(doc)
        s.commit()
        s.refresh(doc)
        out = serialize(doc)
    # 索引失败不算上传失败（embedding 未配置/服务异常只记录到文档状态）
    try:
        rag_service.index_document(doc.id)
    except Exception as e:  # noqa: BLE001
        with Session(engine) as s2:
            fresh = s2.get(Document, doc.id)
            if fresh:
                fresh.index_error = f"向量化失败：{e}"[:300]
                s2.add(fresh)
                s2.commit()
    with Session(engine) as s:
        fresh = s.get(Document, doc.id)
        if fresh:
            out = serialize(fresh)
    return out


def import_local(path: str, *, category_id: Optional[int] = None, title: Optional[str] = None) -> dict:
    """从本地路径导入（供研究页上传材料复用）。"""
    p = Path(path).expanduser()
    if not p.exists():
        raise ValueError("文件不存在")
    return save_upload(p.name, p.read_bytes(), category_id=category_id, title=title)


def serialize(doc: Document, *, with_text: bool = False) -> dict:
    import json

    mindmap = None
    quiz = []
    if doc.mindmap_json:
        try:
            mindmap = json.loads(doc.mindmap_json)
        except Exception:
            mindmap = None
    if doc.quiz_json:
        try:
            quiz = json.loads(doc.quiz_json)
        except Exception:
            quiz = []
    out = {
        "id": doc.id,
        "title": doc.title,
        "filename": doc.filename,
        "path": doc.path,
        "fmt": doc.fmt,
        "icon": _FMT_ICON.get(doc.fmt, "📄"),
        "size_bytes": doc.size_bytes,
        "size_human": human_size(doc.size_bytes),
        "pages": doc.pages,
        "category_id": doc.category_id,
        "uploaded_at": doc.uploaded_at,
        "analyzed": bool(doc.summary_md or doc.mindmap_json or doc.quiz_json),
        "summary_md": doc.summary_md,
        "mindmap": mindmap,
        "quiz": quiz,
        "analyzed_at": doc.analyzed_at,
        "analysis_error": doc.analysis_error,
        "indexed": doc.indexed,
        "indexed_at": doc.indexed_at,
        "chunk_count": doc.chunk_count,
        "index_error": doc.index_error,
        "text_chars": doc.text_chars,
    }
    if with_text:
        out["text"] = doc.text_content
    return out


def list_documents(*, fmt: Optional[str] = None, analyzed: Optional[bool] = None,
                   category_id: Optional[int] = None) -> list[dict]:
    with Session(engine) as s:
        docs = s.exec(select(Document).order_by(Document.uploaded_at.desc())).all()
        out = []
        for d in docs:
            if fmt and d.fmt != fmt:
                continue
            if analyzed is not None and bool(d.summary_md or d.mindmap_json or d.quiz_json) != analyzed:
                continue
            if category_id is not None and d.category_id != category_id:
                continue
            out.append(serialize(d))
        return out


def get_document(doc_id: int, *, with_text: bool = False) -> dict:
    with Session(engine) as s:
        doc = s.get(Document, doc_id)
        if not doc:
            raise KeyError("文档不存在")
        return serialize(doc, with_text=with_text)


_UNSET = object()


def update_document(doc_id: int, *, title: Optional[str] = None,
                    category_id: Any = _UNSET) -> dict:
    with Session(engine) as s:
        doc = s.get(Document, doc_id)
        if not doc:
            raise KeyError("文档不存在")
        if title is not None and title.strip():
            doc.title = title.strip()
        if category_id is not _UNSET:
            doc.category_id = category_id or None  # None/0 = 移出分类
        s.add(doc)
        s.commit()
        s.refresh(doc)
        return serialize(doc)


def delete_document(doc_id: int) -> None:
    from app.services import rag_service

    with Session(engine) as s:
        doc = s.get(Document, doc_id)
        if not doc:
            raise KeyError("文档不存在")
        s.delete(doc)
        s.commit()
    rag_service.remove_document(doc_id)
    (DATA_DIR / doc.path).unlink(missing_ok=True)


def doc_abs_path(doc: Document) -> Path:
    return DATA_DIR / doc.path


def human_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.1f} {unit}" if unit != "B" else f"{n} B"
        n /= 1024
    return f"{n} B"


def reveal_in_finder(doc_id: int) -> None:
    """在访达中显示（macOS）。"""
    import subprocess

    with Session(engine) as s:
        doc = s.get(Document, doc_id)
        if not doc:
            raise KeyError("文档不存在")
        path = doc_abs_path(doc)
    subprocess.Popen(["open", "-R", str(path)])


def build_preview_html(doc: Document) -> str:
    """docx/pptx 服务端转简易 HTML（原文查看用）。"""
    path = doc_abs_path(doc)
    if doc.fmt == "docx":
        from docx import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        import html as _html

        d = DocxDocument(str(path))
        parts = []
        for block in d.iter_inner_content():
            if isinstance(block, Paragraph):
                t = block.text.strip()
                if not t:
                    continue
                style = (block.style.name or "") if block.style else ""
                if style.startswith("Heading"):
                    try:
                        level = min(6, max(1, int(style.split()[-1])))
                    except Exception:
                        level = 2
                    parts.append(f"<h{level}>{_html.escape(t)}</h{level}>")
                else:
                    parts.append(f"<p>{_html.escape(t)}</p>")
            elif isinstance(block, Table):
                rows = []
                for ri, row in enumerate(block.rows):
                    cells = "".join(
                        f"<{'th' if ri == 0 else 'td'}>{_html.escape(c.text.strip())}</{'th' if ri == 0 else 'td'}>"
                        for c in row.cells)
                    rows.append(f"<tr>{cells}</tr>")
                parts.append(f"<table>{''.join(rows)}</table>")
        return "\n".join(parts)
    if doc.fmt == "pptx":
        from pptx import Presentation

        import html as _html

        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = "".join(run.text for run in para.runs).strip()
                        if t:
                            texts.append(f"<p>{_html.escape(t)}</p>")
            parts.append(f'<div class="slide"><div class="slide-no">第 {i} 页</div>{"".join(texts) or "<p class=dim>（空页）</p>"}</div>')
        return "\n".join(parts)
    raise ValueError(f"格式 {doc.fmt} 不需要服务端预览")


def clean_filename(s: str) -> str:
    return re.sub(r"[\\/:*?\"<>|]", "_", s)[:80]
